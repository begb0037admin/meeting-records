import json
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT / "hs_workflow_copy.xlsx"
OUTPUT = ROOT / "hs-items.json"

UPDATE_START = re.compile(
    r"^\s*(\d{1,2}/\d{1,2}/\d{2,4})\s*[-:â€“–—]*\s*(.*)$"
)


def clean_text(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def format_date(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%d %b %Y")
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return str(value).strip() or None


def format_osm(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def parse_updates(value):
    if value is None:
        return []

    entries = []
    current = None
    for raw_line in str(value).splitlines():
        line = raw_line.strip()
        if not line:
            continue

        match = UPDATE_START.match(raw_line)
        if match:
            current = [match.group(1), match.group(2).strip()]
            entries.append(current)
        elif current is None:
            entries.append(["undated", line])
            current = entries[-1]
        else:
            current[1] = f"{current[1]} {line}".strip()

    return entries


def main():
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    try:
        sheet = workbook["Backlog Items"]
        header_row = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))
        columns = {}
        for index, header in enumerate(header_row):
            if header is None:
                continue
            name = str(header).strip()
            if name and name not in columns:
                columns[name] = index

        def value(row, header):
            index = columns[header]
            return row[index] if index < len(row) else None

        items = {}
        for row in sheet.iter_rows(min_row=4, values_only=True):
            ref = clean_text(value(row, "Ref Num"))
            if not ref:
                continue
            key = ref
            if key in items:
                key = f"{ref}_dup"

            items[key] = {
                "item_name": clean_text(value(row, "Item Name")),
                "system": clean_text(value(row, "System")),
                "type": clean_text(value(row, "Type of Request")),
                "moscow": clean_text(value(row, "MoSCoW rating:")),
                "tshirt": clean_text(value(row, "T Shirt Size")),
                "owner": clean_text(value(row, "Owner")),
                "start": format_date(value(row, "Start Date")),
                "exp_delivery": format_date(value(row, "Expected Delivery Date")),
                "delivered": format_date(value(row, "Delivered Date")),
                "osm": format_osm(value(row, "OSM Reference")),
                "detail": clean_text(value(row, "Detail")),
                "benefit": clean_text(value(row, "Expected benefit")),
                "status": clean_text(value(row, "Status")),
                "updates": parse_updates(value(row, "Comment/Updates")),
            }
    finally:
        workbook.close()

    with OUTPUT.open("w", encoding="utf-8") as handle:
        json.dump(items, handle, ensure_ascii=False, indent=2)
        handle.write("\n")

    counts = {}
    for item in items.values():
        status = item.get("status") or "(blank)"
        counts[status] = counts.get(status, 0) + 1
    print(f"items: {len(items)}")
    print("status counts:")
    for status in sorted(counts):
        print(f"  {status}: {counts[status]}")


if __name__ == "__main__":
    main()
